"""
Script simples para gerar Excel pronto para usar
Execute: python criar_excel.py
Resultado: dados_cnpj.xlsx
"""

import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
except ImportError:
    print("❌ Instalando dependências...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side


def criar_excel():
    """Cria arquivo Excel modelo para consulta de CNPJs"""
    
    print("📝 Gerando Excel...")
    
    # Criar workbook
    wb = Workbook()
    
    # ===== ABA 1: INSTRUÇÕES =====
    ws_instrucoes = wb.active
    ws_instrucoes.title = "📖 INSTRUÇÕES"
    
    # Título
    ws_instrucoes.merge_cells('A1:D1')
    titulo = ws_instrucoes['A1']
    titulo.value = "🏢 CONSULTOR SIMPLES NACIONAL"
    titulo.font = Font(size=16, bold=True, color="FFFFFF")
    titulo.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws_instrucoes.row_dimensions[1].height = 35
    
    # Instruções
    linha = 3
    
    # Seção 1
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "✅ PASSO 1: Vá para a aba 'DADOS'"
    cell.font = Font(size=12, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws_instrucoes.row_dimensions[linha].height = 25
    linha += 1
    
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "Coloque seus CNPJs na coluna A (um CNPJ por linha)"
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 1
    
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "Formatos aceitos: 11.222.333/0001-81  ou  11222333000181"
    cell.font = Font(italic=True)
    cell.alignment = Alignment(wrap_text=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 2
    
    # Seção 2
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "✅ PASSO 2: Salve este arquivo"
    cell.font = Font(size=12, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws_instrucoes.row_dimensions[linha].height = 25
    linha += 1
    
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "Use Ctrl+S para salvar"
    cell.alignment = Alignment(wrap_text=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 2
    
    # Seção 3
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "✅ PASSO 3: Execute o script Python"
    cell.font = Font(size=12, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws_instrucoes.row_dimensions[linha].height = 25
    linha += 1
    
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "Abra o terminal e execute: python consultor_cnpj.py"
    cell.font = Font(name="Courier New", size=10, bold=True)
    cell.alignment = Alignment(wrap_text=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 2
    
    # Seção 4
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "✅ PASSO 4: Veja o resultado"
    cell.font = Font(size=12, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws_instrucoes.row_dimensions[linha].height = 25
    linha += 1
    
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "Um novo arquivo será gerado: relatorio_simples_nacional.xlsx"
    cell.alignment = Alignment(wrap_text=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 2
    
    # Legenda de cores
    ws_instrucoes.merge_cells(f'A{linha}:D{linha}')
    cell = ws_instrucoes[f'A{linha}']
    cell.value = "🎨 LEGENDA DE CORES NO RESULTADO"
    cell.font = Font(size=12, bold=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 1
    
    # Cor verde
    cell_verde = ws_instrucoes[f'A{linha}']
    cell_verde.value = "VERDE"
    cell_verde.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    cell_verde.font = Font(bold=True, color="006100")
    
    ws_instrucoes.merge_cells(f'B{linha}:D{linha}')
    cell_desc = ws_instrucoes[f'B{linha}']
    cell_desc.value = "Optante pelo Simples Nacional ✅"
    cell_desc.alignment = Alignment(wrap_text=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 1
    
    # Cor vermelha
    cell_vermelho = ws_instrucoes[f'A{linha}']
    cell_vermelho.value = "VERMELHO"
    cell_vermelho.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    cell_vermelho.font = Font(bold=True, color="9C0006")
    
    ws_instrucoes.merge_cells(f'B{linha}:D{linha}')
    cell_desc = ws_instrucoes[f'B{linha}']
    cell_desc.value = "NÃO é optante pelo Simples Nacional ❌"
    cell_desc.alignment = Alignment(wrap_text=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    linha += 1
    
    # Cor amarela
    cell_amarelo = ws_instrucoes[f'A{linha}']
    cell_amarelo.value = "AMARELO"
    cell_amarelo.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    cell_amarelo.font = Font(bold=True, color="9C6500")
    
    ws_instrucoes.merge_cells(f'B{linha}:D{linha}')
    cell_desc = ws_instrucoes[f'B{linha}']
    cell_desc.value = "Erro (CNPJ inválido, não encontrado, etc) ⚠️"
    cell_desc.alignment = Alignment(wrap_text=True)
    ws_instrucoes.row_dimensions[linha].height = 20
    
    # Configurar largura
    ws_instrucoes.column_dimensions['A'].width = 20
    ws_instrucoes.column_dimensions['B'].width = 50
    ws_instrucoes.column_dimensions['C'].width = 15
    ws_instrucoes.column_dimensions['D'].width = 15
    
    # ===== ABA 2: DADOS =====
    ws_dados = wb.create_sheet("📊 DADOS")
    
    # Estilos
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Header
    ws_dados.cell(row=1, column=1).value = "CNPJ"
    ws_dados.cell(row=1, column=1).fill = header_fill
    ws_dados.cell(row=1, column=1).font = header_font
    ws_dados.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws_dados.cell(row=1, column=1).border = border
    
    ws_dados.cell(row=1, column=2).value = "Observações (opcional)"
    ws_dados.cell(row=1, column=2).fill = header_fill
    ws_dados.cell(row=1, column=2).font = header_font
    ws_dados.cell(row=1, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws_dados.cell(row=1, column=2).border = border
    
    ws_dados.row_dimensions[1].height = 25
    
    # Preencher linhas vazias (100 linhas para dados)
    for i in range(2, 102):
        ws_dados.cell(row=i, column=1).border = border
        ws_dados.cell(row=i, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws_dados.cell(row=i, column=2).border = border
        ws_dados.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    
    # Largura das colunas
    ws_dados.column_dimensions['A'].width = 25
    ws_dados.column_dimensions['B'].width = 50
    
    # Salvar arquivo
    wb.save('dados_cnpj.xlsx')
    
    print("✅ Excel criado com sucesso!")
    print("📁 Arquivo: dados_cnpj.xlsx")
    print("📍 Localização: Na mesma pasta deste script")
    print("\n🚀 Próximo passo:")
    print("   1. Abra 'dados_cnpj.xlsx'")
    print("   2. Vá para a aba '📊 DADOS'")
    print("   3. Coloque seus CNPJs na coluna A")
    print("   4. Salve o arquivo (Ctrl+S)")
    print("   5. Execute: python consultor_cnpj.py")
    print("\n✨ Pronto! Aguarde o resultado!")


if __name__ == "__main__":
    try:
        criar_excel()
    except Exception as e:
        print(f"❌ Erro: {e}")
        print("\nTente instalar manualmente:")
        print("pip install openpyxl")
