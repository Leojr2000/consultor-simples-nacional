from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# Criando o workbook
wb = Workbook()
ws = wb.active
ws.title = "CNPJs para Consultar"

# Estilos
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=12)
example_fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Header
ws.cell(row=1, column=1).value = "CNPJ"
ws.cell(row=1, column=1).fill = header_fill
ws.cell(row=1, column=1).font = header_font
ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
ws.cell(row=1, column=1).border = border

ws.cell(row=1, column=2).value = "Observações"
ws.cell(row=1, column=2).fill = header_fill
ws.cell(row=1, column=2).font = header_font
ws.cell(row=1, column=2).alignment = Alignment(horizontal="center", vertical="center")
ws.cell(row=1, column=2).border = border

# Exemplos
exemplos = [
    ("11.222.333/0001-81", "Exemplo com formatação"),
    ("11222333000181", "Exemplo sem formatação"),
    ("", "Cole seus CNPJs aqui"),
    ("", ""),
    ("", ""),
]

for i, (cnpj, obs) in enumerate(exemplos, 2):
    ws.cell(row=i, column=1).value = cnpj
    ws.cell(row=i, column=1).fill = example_fill
    ws.cell(row=i, column=1).border = border
    ws.cell(row=i, column=1).alignment = Alignment(horizontal="center")
    
    ws.cell(row=i, column=2).value = obs
    ws.cell(row=i, column=2).fill = example_fill
    ws.cell(row=i, column=2).border = border

# Adicionar mais linhas vazias para preenchimento
for i in range(7, 52):
    ws.cell(row=i, column=1).border = border
    ws.cell(row=i, column=2).border = border
    ws.cell(row=i, column=1).alignment = Alignment(horizontal="center")

# Largura das colunas
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 40

# Altura do header
ws.row_dimensions[1].height = 25

# Instrução na aba "Instruções"
wb_info = wb.create_sheet("Como Usar", 0)
ws_info = wb_info

ws_info.merge_cells('A1:D1')
cell = ws_info['A1']
cell.value = "📋 CONSULTOR SIMPLES NACIONAL"
cell.font = Font(size=16, bold=True, color="FFFFFF")
cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
cell.alignment = Alignment(horizontal="center", vertical="center")
ws_info.row_dimensions[1].height = 30

ws_info.merge_cells('A3:D3')
ws_info['A3'].value = "✅ INSTRUÇÕES DE USO"
ws_info['A3'].font = Font(size=12, bold=True)

instrucoes = [
    ("", ""),
    ("1️⃣ PASSO 1 - PREPARAR OS DADOS", "Na aba 'CNPJs para Consultar', coloque seus CNPJs na coluna A"),
    ("", "Cada CNPJ deve estar em uma linha diferente"),
    ("", "Aceita formatação: 11.222.333/0001-81 ou sem: 11222333000181"),
    ("", ""),
    ("2️⃣ PASSO 2 - SALVE ESTE ARQUIVO", "Use Ctrl+S para salvar o arquivo"),
    ("", "Nome sugerido: dados_cnpj.xlsx"),
    ("", ""),
    ("3️⃣ PASSO 3 - EXECUTE O SCRIPT", "Abra o terminal/prompt na pasta do projeto"),
    ("", "Digite: python consultor_cnpj.py"),
    ("", "Aguarde a consulta terminar"),
    ("", ""),
    ("4️⃣ PASSO 4 - VEJA OS RESULTADOS", "Será gerado um arquivo: relatorio_simples_nacional.xlsx"),
    ("", "Abra e veja as cores:"),
    ("🟢 VERDE", "= Optante pelo Simples Nacional"),
    ("🔴 VERMELHO", "= NÃO é optante"),
    ("🟡 AMARELO", "= Erro (CNPJ inválido, não encontrado)"),
    ("", ""),
    ("❓ DÚVIDAS FREQUENTES", ""),
    ("Qual formato de CNPJ?", "Tanto com formatação (11.222.333/0001-81) quanto sem (11222333000181)"),
    ("Quantos CNPJs posso consultar?", "Sem limite! 10, 100, 1000... todos são consultados"),
    ("Custa algo?", "NÃO! As APIs utilizadas são 100% gratuitas"),
    ("Quanto tempo leva?", "Depende da quantidade: ~5-10 segundos por CNPJ"),
    ("", ""),
]

row = 5
for titulo, descricao in instrucoes:
    if titulo:
        cell_titulo = ws_info.cell(row=row, column=1)
        cell_titulo.value = titulo
        cell_titulo.font = Font(bold=True, size=11, color="1F4E78")
        
        cell_desc = ws_info.cell(row=row, column=2)
        cell_desc.value = descricao
        cell_desc.alignment = Alignment(wrap_text=True, vertical="top")
        
        ws_info.merge_cells(f'B{row}:D{row}')
        ws_info.row_dimensions[row].height = None
    row += 1

ws_info.column_dimensions['A'].width = 30
ws_info.column_dimensions['B'].width = 50

# Salvando o arquivo
wb.save('dados_cnpj.xlsx')
print("✅ Arquivo 'dados_cnpj.xlsx' gerado com sucesso!")
print("📁 Está pronto para download!")
